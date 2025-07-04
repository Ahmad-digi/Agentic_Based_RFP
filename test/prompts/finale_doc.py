from docx import Document
from markdown import markdown
from bs4 import BeautifulSoup
import io
def append_markdown_to_docx(markdown_text):
    buffer = io.BytesIO()
    doc = Document()

    html = markdown(markdown_text)
    soup = BeautifulSoup(html, "html.parser")
    for element in soup.contents:
        if element.name == "h1":
            doc.add_heading(element.text, level=1)
        elif element.name == "h2":
            doc.add_heading(element.text, level=2)
        elif element.name == "h3":
            doc.add_heading(element.text, level=3)
        elif element.name == "ul":
            for li in element.find_all("li"):
                doc.add_paragraph(li.text, style='ListBullet')
        elif element.name == "ol":
            for li in element.find_all("li"):
                doc.add_paragraph(li.text, style='ListNumber')
        elif element.name == "p":
            doc.add_paragraph(element.text)
        elif element.name == "pre":
            code_text = element.text
            doc.add_paragraph(code_text, style='Intense Quote')
        elif element.text and element.text.strip():
            doc.add_paragraph(element.text)

    doc.save(buffer)
    buffer.seek(0)
    return buffer